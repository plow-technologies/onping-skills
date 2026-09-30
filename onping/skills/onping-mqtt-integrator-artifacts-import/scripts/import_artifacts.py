# /// script
# requires-python = ">=3.10"
# dependencies = ["requests", "openpyxl"]
# ///
"""Upload an artifacts spreadsheet to mqtt-json-integrator's created-object record.

Route: POST /mqtt/json/integrator/{serial}/artifacts/import
Body:  multipart form, f1 = XLSX bytes

THIS CREATES NOTHING. The handler routes to postArtifacts, a repsert of the
integrator's OWN record store. It does not touch OnPing or the mqtt-json driver.
The route that actually creates objects is POST .../artifacts, wrapped by
onping-mqtt-integrator-create — a genuinely confusing pair of names, so:

    .../artifacts/import   -> record only, creates nothing   (this skill)
    .../artifacts          -> five-stage CREATE of real objects

WHAT IT IS ACTUALLY FOR: SUPPRESSION. The create pipeline skips anything already
recorded as created. So importing a record of objects that already exist stops
the integrator from creating them again — the documented use is preventing
Creatable Objects from being turned into Created Objects
(data-lifetime-and-uniqueness-rules.md, "Created Objects"). It is also how you
restore a record after a --target stored deletion.

REPSERT, NOT REPLACE. Unlike the rules import, this one is keyed: entries are
upserted into the stored sets by their unique identifiers. Rows absent from the
sheet are NOT deleted. Use onping-mqtt-integrator-delete --target stored to
remove records.

COLUMN 3 IS TRUSTED VERBATIM. "Location ID Ref" becomes both storedLocationIdRef
and every PID's storedPidLocationIdRef, with no verification that the refId
exists in OnPing. A wrong number produces a record pointing at the wrong
location, or at nothing.

THE ROUND-TRIP IS LOSSY (Artifacts/ImportExport.hs): localParameterTime is
dropped and unrecognized values collapse to Double 0.0 / "0.0". Re-importing an
exported sheet therefore DEGRADES the record. Back up with
onping-mqtt-integrator-artifacts-export --json (lossless) first.

THE f1 FIELD-NAME GOTCHA: the handler's `fileAFormReq "File"` gives a LABEL, not
a field name; Yesod names fields positionally, so it is `f1`. Posting "File"
yields 400 FormFailure.

MUTATES OnPing (well, the integrator's store) only with --yes.

Source of truth (re-verify if these drift):
  - Handler: onping/Handler/MqttJsonIntegrator/Service.hs
  - Rebuild: onping/Handler/MqttJsonIntegrator/Artifacts/ImportExport.hs
             (reconstructArtifacts)
  - Schema:  _mqtt_integrator_routes/rules_sheet.py (ARTIFACTS_HEADERS)
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))  # the skills root
from _mqtt_integrator_routes import rules_sheet
from _mqtt_integrator_routes.integrator_http import (
    _fail,
    get_json,
    post_multipart_xlsx,
    report_write,
)
from _mqtt_integrator_routes.routes import endpoint

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - uv installs the dep
    print("openpyxl is required (declared in the uv script header)", file=sys.stderr)
    sys.exit(1)


def read_local_sheet(path: Path) -> tuple[list, list[list]]:
    try:
        wb = load_workbook(path, data_only=True)
    except Exception as e:
        _fail(f"Could not read {path}: {e}")
    if rules_sheet.SHEET_NAME not in wb.sheetnames:
        _fail(
            f"{path} has no {rules_sheet.SHEET_NAME!r} worksheet (found "
            f"{wb.sheetnames}). The server reads only "
            f"{rules_sheet.SHEET_NAME!r}."
        )
    ws = wb[rules_sheet.SHEET_NAME]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    if not rows:
        _fail(f"{path} is empty.")
    header = list(rows[rules_sheet.HEADER_ROW - 1])
    data = rows[rules_sheet.FIRST_DATA_ROW - 1 :]

    def is_blank(row: list) -> bool:
        return all(v is None or (isinstance(v, str) and not v.strip()) for v in row)

    while data and is_blank(data[-1]):
        data.pop()
    return header, data


def main() -> None:
    p = argparse.ArgumentParser(
        description=(
            "Upload an artifacts XLSX to the mqtt-json-integrator created-object "
            "record. Creates NOTHING in OnPing — it repserts the integrator's own "
            "record, which suppresses future creation. Mutates only with --yes."
        ),
    )
    p.add_argument("access_token", help="OnPing OAuth bearer token")
    p.add_argument("serial", help="LJSerial of the Lumberjack running the integrator")
    p.add_argument("spreadsheet", help="path to the artifacts XLSX to upload")
    p.add_argument(
        "--skip-live-check",
        action="store_true",
        help="do not fetch the live record for the new/overwritten comparison",
    )
    p.add_argument(
        "--dry-run",
        action="store_true",
        help="validate and preview; never POST (wins over --yes)",
    )
    p.add_argument("--yes", action="store_true", help="actually perform the upload")
    args = p.parse_args()

    path = Path(args.spreadsheet)
    if not path.is_file():
        _fail(f"No such file: {path}")

    header, rows = read_local_sheet(path)

    expected = rules_sheet.ARTIFACTS_HEADERS
    got = [str(h).strip() if h is not None else "" for h in header][: len(expected)]
    if got != expected:
        print(
            f"Header row does not match the expected artifacts-sheet schema.\n"
            f"  expected: {expected}\n  got:      {got}",
            file=sys.stderr,
        )
        if got == rules_sheet.RULES_HEADERS[: len(got)]:
            print(
                "\nThis looks like a RULES sheet, not an artifacts sheet. The two "
                "are both 12 columns but mean different things — you want "
                "onping-mqtt-integrator-rules-import.",
                file=sys.stderr,
            )
        _fail("Refusing to upload a sheet whose headers do not match.")

    if not rows:
        _fail("The sheet has no data rows; there is nothing to record.")

    errors: list[str] = []
    location_only = 0
    pid_rows = 0
    locations: dict[str, tuple[str, object]] = {}  # key -> (name, refId)
    pid_keys: set[tuple[str, str, str]] = set()

    for offset, row in enumerate(rows):
        row_number = rules_sheet.FIRST_DATA_ROW + offset
        row_errors, is_loc_only = rules_sheet.validate_row(
            row, flavor="artifacts", row_number=row_number
        )
        errors.extend(row_errors)

        cells = list(row) + [None] * (12 - len(row))
        key = str(cells[0]).strip() if cells[0] is not None else ""
        name = str(cells[1]).strip() if cells[1] is not None else ""
        refid = cells[2]

        # Column 3 must be an integer location refId; the server casts it
        # without checking, so catch a bad value here.
        if isinstance(refid, str):
            try:
                refid = int(refid.strip())
            except ValueError:
                errors.append(
                    f"row {row_number}, column 3 (Location ID Ref): {cells[2]!r} is "
                    f"not an integer. It becomes the OnPing location refId verbatim."
                )
        elif isinstance(refid, float) and not refid.is_integer():
            errors.append(
                f"row {row_number}, column 3 (Location ID Ref): {refid!r} is not a "
                f"whole number."
            )

        if key in locations and locations[key] != (name, refid):
            errors.append(
                f"row {row_number}: location key {key!r} appears with conflicting "
                f"name/refId — {locations[key]} vs {(name, refid)}. The record is "
                f"keyed by column 1, so one of these silently wins."
            )
        locations.setdefault(key, (name, refid))

        if is_loc_only:
            location_only += 1
        else:
            pid_rows += 1
            pid_keys.add(
                (
                    key,
                    str(cells[4]).strip() if cells[4] is not None else "",
                    str(cells[5]).strip() if cells[5] is not None else "",
                )
            )

    print(
        f"Sheet: {path}\n"
        f"  {len(rows)} data row(s): {pid_rows} PID row(s), "
        f"{location_only} location-only\n"
        f"  {len(locations)} distinct location(s), {len(pid_keys)} distinct PID(s)\n"
    )

    if pid_rows != len(pid_keys):
        print(
            f"WARNING: {pid_rows - len(pid_keys)} PID row(s) duplicate an existing "
            f"(location, topic, value-selector) key. PIDs are stored in a SET keyed "
            f"on those fields, so duplicates collapse silently.",
            file=sys.stderr,
        )

    if errors:
        print(f"\n{len(errors)} problem(s) with this sheet:", file=sys.stderr)
        for e in errors:
            print(f"  - {e}", file=sys.stderr)
        _fail("Refusing to upload an invalid sheet.")

    # ── compare against the live record ─────────────────────────────────────
    if args.skip_live_check:
        print("Skipping the live-record comparison (--skip-live-check).\n")
    else:
        live = get_json(
            args.access_token, endpoint("get_artifacts"), serial=args.serial
        )
        live_locs = {
            (l.get("storedLocationUniqueIdentifier") or {}).get(
                "unLocationUniqueIdentifier"
            )
            for l in (live.get("storedLocations") or [])
        }
        live_pid_count = len(live.get("storedPids") or [])
        new_locs = [k for k in locations if k not in live_locs]
        overwritten = [k for k in locations if k in live_locs]
        print(
            f"Live record on LJ {args.serial}: {len(live_locs)} location(s), "
            f"{live_pid_count} PID(s)\n"
            f"  new locations         : {len(new_locs)}\n"
            f"  overwritten locations : {len(overwritten)}\n"
            f"  untouched (kept)      : {len(live_locs - set(locations))}   "
            f"(this is a repsert — absent rows are NOT deleted)\n"
        )

    print(
        "This import creates NOTHING in OnPing or the mqtt-json driver. Its effect "
        "is on the integrator's record — and therefore on SUPPRESSION: any object "
        "recorded here will be skipped by onping-mqtt-integrator-create."
    )
    print(
        "\nNote the export that produced this sheet is LOSSY: PID timestamps are "
        "dropped and unrecognized values collapse to 0.0. Re-importing an exported "
        "sheet degrades the record."
    )

    if args.dry_run or not args.yes:
        why = "--dry-run" if args.dry_run else "no --yes"
        print(f"\nNot uploaded ({why}). Re-run with --yes to apply.")
        print(
            "Back up losslessly first: onping-mqtt-integrator-artifacts-export "
            f"<token> {args.serial} --json"
        )
        return

    resp = post_multipart_xlsx(
        args.access_token,
        endpoint("import_artifacts"),
        path.read_bytes(),
        filename=path.name,
        serial=args.serial,
    )
    report_write(resp, what=f"artifacts import for LJ {args.serial}")
    print(
        f"\nRecorded {len(locations)} location(s) and {len(pid_keys)} PID(s) on "
        f"LJ {args.serial}."
    )


if __name__ == "__main__":
    main()
