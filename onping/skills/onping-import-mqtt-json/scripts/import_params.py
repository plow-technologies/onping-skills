# /// script
# requires-python = ">=3.10"
# dependencies = ["requests", "openpyxl"]
# ///
"""Upload (round-trip) an mqtt-json parameter spreadsheet to OnPing.

This is a **bulk update**, not an author-from-scratch tool. `POST
/mqtt/json/param/import` is a multipart form (`File` + `Location` int) that
reconciles the location's ENTIRE parameter set from the sheet: the server
rejects the import unless every existing PID appears in the file
(`existingPids subset-of importedPids`). The only safe workflow is therefore:

    export current (onping-export-mqtt-json) -> edit rows -> re-upload here

MUTATES live OnPing state only with --yes. Without --yes (or with --dry-run)
the skill validates the file against the live location and prints a preview,
never POSTing.

Source of truth (re-verify if these drift):
  - Handler:  onping/Handler/MQTT/JSON/Service.hs (postImportMqttJsonParametersR)
  - Parser:   onping/Handler/MQTT/JSON/ImportExport.hs (ParameterInfoExportable)
  - Export:   GET /mqtt/json/param/export/{loc}/{filename}  (read side, reused here)
"""

from __future__ import annotations

import argparse
import io
import os
import sys

import requests

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - uv installs the dep
    print("openpyxl is required (declared in the uv script header)", file=sys.stderr)
    sys.exit(1)

BASE_URL = "https://onping.plowtech.net"
TIMEOUT_SECONDS = 60

IMPORT_PATH = "/mqtt/json/param/import"
EXPORT_PATH = "/mqtt/json/param/export/{loc}/{filename}"

XLSX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)

# The 8 columns, in order, per ImportExport.hs `xlsxSheetHeaders`.
EXPECTED_HEADERS = [
    "PID",
    "Description",
    "Topic",
    "Value Selector",
    "Type",
    "Time Format",
    "Time Selector",
    "Writeable",
]

# `Type` column legal values (ParameterValueTypeAssigned rendering). Empty is
# allowed; the column is parsed but IGNORED on import (stored value type is
# preserved for existing PIDs), so we validate leniently.
VALID_TYPES = {
    "",
    "Boolean",
    "Double",
    "Utf8 Text 24",
    "Utf8 Text 40",
    "Utf8 Text 184",
}


def _cell_str(value) -> str:
    """Normalize an openpyxl cell value to a trimmed string."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value).strip()


def _parse_bool_cell(raw: str) -> bool | None:
    """Match the Haskell XlsxCell Bool instance: TRUE/FALSE (any case) or 1/0."""
    low = raw.strip().lower()
    if low in ("true", "1", "1.0"):
        return True
    if low in ("false", "0", "0.0"):
        return False
    return None


def _load_rows(xlsx_bytes: bytes) -> tuple[list[str], list[list[str]]]:
    """Return (header_row, data_rows) as strings. Data starts at sheet row 2."""
    wb = load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return [], []
    header = [_cell_str(c) for c in rows[0]][: len(EXPECTED_HEADERS)]
    data = []
    for raw_row in rows[1:]:
        cells = [_cell_str(c) for c in raw_row]
        # pad/truncate to 8 columns
        cells = (cells + [""] * len(EXPECTED_HEADERS))[: len(EXPECTED_HEADERS)]
        if all(c == "" for c in cells):
            continue  # skip fully blank trailing rows
        data.append(cells)
    return header, data


def _validate(header: list[str], data: list[list[str]]) -> list[str]:
    """Return a list of human-readable validation errors (empty == valid)."""
    errors: list[str] = []
    if header != EXPECTED_HEADERS:
        errors.append(
            "Header row does not match the 8 expected columns.\n"
            f"  expected: {EXPECTED_HEADERS}\n"
            f"  found:    {header}"
        )
        # Header mismatch means column positions are untrustworthy; stop here.
        return errors

    for i, row in enumerate(data):
        sheet_row = i + 2  # 1-based, header is row 1
        pid, _desc, _topic, _valsel, vtype, tfmt, _timesel, write = row

        if pid != "":
            # PID cell, when present, must be an integer.
            try:
                int(float(pid))
            except ValueError:
                errors.append(f"row {sheet_row}: PID '{pid}' is not an integer")

        if vtype not in VALID_TYPES:
            errors.append(
                f"row {sheet_row}: Type '{vtype}' is not one of "
                f"{sorted(VALID_TYPES - {''})} (or blank)"
            )

        if _parse_bool_cell(write) is None:
            errors.append(
                f"row {sheet_row}: Writeable '{write}' is not a boolean "
                "(TRUE/FALSE or 1/0)"
            )

        # Time Format: ISO, LumberjackTime, or any non-empty custom strftime.
        if tfmt == "":
            errors.append(
                f"row {sheet_row}: Time Format is empty "
                "(expected ISO, LumberjackTime, or a custom strftime string)"
            )
    return errors


def _sheet_pids(data: list[str]) -> set[int]:
    pids: set[int] = set()
    for row in data:
        pid = row[0]
        if pid != "":
            try:
                pids.add(int(float(pid)))
            except ValueError:
                pass
    return pids


def _fetch_existing_pids(token: str, loc: int) -> set[int]:
    """Fetch the location's current params via the export route; return PIDs.

    Uses the exact same 8-column XLSX projection the server round-trips, so the
    local `existingPids subset-of importedPids` check matches the server's.
    """
    url = BASE_URL + EXPORT_PATH.format(loc=loc, filename="current.xlsx")
    headers = {"Authorization": f"Bearer {token}"}
    try:
        resp = requests.get(
            url, headers=headers, allow_redirects=True, timeout=TIMEOUT_SECONDS
        )
    except requests.RequestException as e:
        print(f"Failed to fetch existing params: {e}", file=sys.stderr)
        sys.exit(1)

    if not (200 <= resp.status_code < 300):
        print(
            f"HTTP {resp.status_code} fetching existing params from {url}\n"
            f"{resp.text[:500]}",
            file=sys.stderr,
        )
        sys.exit(1)

    body = resp.content
    if body.lstrip()[:5].lower() in (b"<!doc", b"<html"):
        print(
            "Expected XLSX when fetching existing params but got HTML — "
            "the access token may be expired.",
            file=sys.stderr,
        )
        sys.exit(1)

    _header, data = _load_rows(body)
    return _sheet_pids(data)


def _upload(token: str, loc: int, path: str, xlsx_bytes: bytes) -> None:
    url = BASE_URL + IMPORT_PATH
    headers = {"Authorization": f"Bearer {token}"}
    # The handler uses `runFormPostNoToken $ renderDivs $ (,) <$> fileAFormReq
    # "File" <*> areq intField "Location"`. Those strings are LABELS; Yesod's
    # renderDivs auto-generates the actual field NAMES `f1`, `f2` in field
    # order. The frontend (OnpingFetch_ImportParameters.res) posts exactly
    # `f1` = file blob, `f2` = location int. Sending "File"/"Location" yields a
    # 400 "FormFailure". Field order: file first, then location.
    files = {"f1": (os.path.basename(path), xlsx_bytes, XLSX_CONTENT_TYPE)}
    data = {"f2": str(loc)}
    try:
        resp = requests.post(
            url,
            headers=headers,
            files=files,
            data=data,
            allow_redirects=True,
            timeout=TIMEOUT_SECONDS,
        )
    except requests.RequestException as e:
        print(f"Upload request failed: {e}", file=sys.stderr)
        sys.exit(1)

    if 200 <= resp.status_code < 300:
        print(f"OK — imported spreadsheet for location {loc}.")
        # The handler returns the updated [(PID, ParameterInfo)] list.
        try:
            payload = resp.json()
            n = len(payload) if isinstance(payload, list) else "?"
            print(f"Server acknowledged {n} parameters for the location.")
        except ValueError:
            pass
        return

    # Surface the server error verbatim (e.g. the missing-PID message).
    print(f"HTTP {resp.status_code} from {url}", file=sys.stderr)
    print(resp.text[:1000], file=sys.stderr)
    sys.exit(1)


def main() -> None:
    p = argparse.ArgumentParser(
        description="Upload/round-trip a mqtt-json parameter spreadsheet to "
        "OnPing (bulk update). MUTATING — requires --yes.",
    )
    p.add_argument("access_token", help="OnPing OAuth bearer token")
    p.add_argument("location_id", type=int, help="location refId (LocationIdRef)")
    p.add_argument("spreadsheet", help="path to the .xlsx to upload")
    p.add_argument(
        "--dry-run",
        action="store_true",
        help="validate + preview only; never POST (wins over --yes)",
    )
    p.add_argument(
        "--yes",
        action="store_true",
        help="perform the upload (POST). Without it, preview only.",
    )
    args = p.parse_args()

    if not os.path.exists(args.spreadsheet):
        print(f"File not found: {args.spreadsheet}", file=sys.stderr)
        sys.exit(1)

    with open(args.spreadsheet, "rb") as f:
        xlsx_bytes = f.read()

    header, data = _load_rows(xlsx_bytes)

    # 1. Format validation.
    errors = _validate(header, data)

    # 2. Classify rows.
    updates = sum(1 for r in data if r[0] != "")
    creates = sum(1 for r in data if r[0] == "")

    # 3. PID-coverage precondition (mirrors the server).
    existing = _fetch_existing_pids(args.access_token, args.location_id)
    imported = _sheet_pids(data)
    missing = sorted(existing - imported)

    # ---- Preview ----
    print(f"Location:      {args.location_id}")
    print(f"Spreadsheet:   {os.path.abspath(args.spreadsheet)}")
    print(f"Rows:          {len(data)}  ({updates} updates, {creates} creates)")
    print(f"Existing PIDs: {len(existing)}")
    if missing:
        print(f"MISSING PIDs:  {missing}")
    else:
        print("Coverage:      OK (all existing PIDs present)")

    if errors:
        print("\nValidation errors:", file=sys.stderr)
        for e in errors:
            print(f"  - {e}", file=sys.stderr)

    # ---- Decide ----
    blocked = bool(errors) or bool(missing)

    if args.dry_run or not args.yes:
        if args.dry_run:
            print("\n[dry-run] not uploading.")
        else:
            print("\nPreview only. Re-run with --yes to upload.")
        if blocked:
            print(
                "NOTE: upload is currently BLOCKED "
                f"({'validation errors' if errors else ''}"
                f"{' and ' if errors and missing else ''}"
                f"{'missing PIDs' if missing else ''}).",
                file=sys.stderr,
            )
        return

    # ---- Upload path (--yes) ----
    if blocked:
        if missing:
            print(
                "\nRefusing to upload: the spreadsheet is missing existing PIDs "
                f"{missing}. Start from a fresh export "
                "(onping-export-mqtt-json) so every current parameter is present.",
                file=sys.stderr,
            )
        if errors:
            print(
                "\nRefusing to upload: fix the validation errors above first.",
                file=sys.stderr,
            )
        sys.exit(1)

    _upload(args.access_token, args.location_id, args.spreadsheet, xlsx_bytes)


if __name__ == "__main__":
    main()
