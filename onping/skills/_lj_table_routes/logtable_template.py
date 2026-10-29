"""Log-table column definitions: the `--column` flag, the export XLSX header, and the v1 schema.

Pure, apart from reading a workbook file. No request is made here.

The export format (`onping-logtable-export`, server `schemaHeadersToCellMapMultipleRowsHeader`)
is a sheet whose first four rows define the columns, one sheet column per log-table column,
in position order (hidden columns are left out):

    row 1   `event` on the event (trigger) column only
    row 2   `name: <columnSchemaName>`
    row 3   `tag: <word>`        a `tagAsText` word: result, unit, dateTime, ...
    row 4   `source: <n>`        a PID, or `source: <n> VPID`
    row 5+  data rows

`read_header` reads it the way the server's `parseXlsxColumnHeaderMultipleRows`
(onping-logtable-app, Handlers/ImportExport.hs) does: `key: value` cells are split
on every ":" and must give exactly two parts (`parseKeyValue`), tag words are
case-sensitive (`parseTagSelector`), and sources are strict (`parsePointerFromText`).
It is stricter in two ways: exactly one column may be marked `event`, and text or
invalid tags are refused, since a log table cannot hold them.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from . import logtable as lt
from .pointers import EVENT_SELECTOR, LOGTABLE_SELECTORS, Pointer, normalize_selector, parse_cell

# tagAsText / parseTagSelector (onping-logtable-types, Onping/LogTable/Xlsx.hs).
TAG_WORDS = {
    "text": "TextTagSelector",
    "pid": "PIDTagSelector",
    "result": "ResultTagSelector",
    "unit": "UnitTagSelector",
    "description": "DescriptionTagSelector",
    "dateTime": "LastUpdateDateTimeTagSelector",
    "date": "LastUpdateDateTagSelector",
    "time": "LastUpdateTimeTagSelector",
    "writeable": "WriteabilityTagSelector",
    "company": "CompanyTagSelector",
    "companyId": "CompanyIdTagSelector",
    "site": "SiteTagSelector",
    "siteId": "SiteIdTagSelector",
    "location": "LocationTagSelector",
    "locationId": "LocationIdTagSelector",
    "invalid": "InvalidTagSelector",
}
WORD_OF = {sel: word for word, sel in TAG_WORDS.items()}
DEFAULT_TAG = "result"
FLAG_FORM = "NAME=SOURCE[:TAG], e.g. 'Casing=500002' or 'Mode=7 VPID:unit'"


@dataclass(frozen=True)
class ColumnSpec:
    name: str
    pointer: Pointer
    selector: str
    event: bool = False


@dataclass
class Header:
    """The columns of an export file, in file order, and how many data rows it holds."""
    columns: list[ColumnSpec]
    data_rows: int = 0
    sheet: str = "Sheet1"
    notes: list[str] = field(default_factory=list)

    @property
    def event(self) -> ColumnSpec:
        return next(c for c in self.columns if c.event)


def tag_word(selector: str) -> str:
    return WORD_OF.get(selector, selector)


def _tag_words() -> str:
    return ", ".join(w for w, s in TAG_WORDS.items() if s in LOGTABLE_SELECTORS)


def selector_for_flag(tag: str) -> str:
    """A `--column` TAG: a tag word in any case, or a selector name (`Result`, `ResultTagSelector`)."""
    t = tag.strip()
    for word, sel in TAG_WORDS.items():
        if word.lower() == t.lower():
            return sel
    try:
        return normalize_selector(t)
    except ValueError:
        raise ValueError(f"unknown tag {t!r}; use one of: {_tag_words()}") from None


def event_spec(event_pid: int, event_name: str = "eventDateTime") -> ColumnSpec:
    return ColumnSpec(event_name, Pointer("PID", event_pid), EVENT_SELECTOR, event=True)


def parse_column_spec(text: str) -> ColumnSpec:
    """Parse one `--column NAME=SOURCE[:TAG]`. Split at the last "=", then at the last ":"."""
    if "=" not in text:
        raise ValueError(f"--column {text!r}: expected {FLAG_FORM}")
    name, rest = text.rsplit("=", 1)
    name = name.strip()
    if not name:
        raise ValueError(f"--column {text!r}: the column name is empty")
    src, tag = rest.rsplit(":", 1) if ":" in rest else (rest, DEFAULT_TAG)
    try:
        pointer = parse_cell(src.strip())
    except ValueError:
        raise ValueError(f"--column {text!r}: source {src.strip()!r} is not '<n>' (a PID) "
                         "or '<n> VPID'") from None
    try:
        selector = selector_for_flag(tag)
    except ValueError as e:
        raise ValueError(f"--column {text!r}: {e}") from None
    if selector not in LOGTABLE_SELECTORS:
        raise ValueError(f"--column {text!r}: a log table cannot hold a {tag_word(selector)!r} "
                         f"column; use one of: {_tag_words()}")
    return ColumnSpec(name, pointer, selector)


# ------------------------------------------------------------------ export header

def _blank(v: Any) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def _key_value(v: Any, row: int, col: int, key: str, errors: list[str]) -> str | None:
    where = f"R{row}C{col}"
    if _blank(v):
        errors.append(f"{where}: missing cell, expected '{key}: <value>'")
        return None
    if not isinstance(v, str):
        errors.append(f"{where}: expected the text '{key}: <value>', found {v!r}")
        return None
    parts = v.split(":")
    if len(parts) != 2:
        hint = " (a column name cannot contain ':' in this format)" if key == "name" else ""
        errors.append(f"{where}: expected exactly one ':' in '{key}: <value>', found {v!r}{hint}")
        return None
    k, value = parts[0].strip(), parts[1].strip()
    if k != key:
        errors.append(f"{where}: expected key '{key}' but found '{k}'")
        return None
    return value


def parse_header_rows(rows: list[list[Any]], *, data_rows: int = 0, sheet: str = "Sheet1") -> Header:
    """Parse rows 1-4 (lists of cell values, row-major) into column specs. Raises ValueError
    listing every problem found."""
    rows = [list(r) for r in rows[:4]] + [[] for _ in range(4 - len(rows[:4]))]
    width = max((i + 1 for r in rows for i, v in enumerate(r) if not _blank(v)), default=0)
    if width == 0:
        raise ValueError("rows 1-4 are empty; expected the export header "
                         "(event / name: / tag: / source:)")

    def cell(r: int, c: int) -> Any:
        return rows[r - 1][c - 1] if c - 1 < len(rows[r - 1]) else None

    errors: list[str] = []
    cols: list[ColumnSpec] = []
    for c in range(1, width + 1):
        marker = cell(1, c)
        is_event = isinstance(marker, str) and marker.strip() == "event"
        name = _key_value(cell(2, c), 2, c, "name", errors)
        tag = _key_value(cell(3, c), 3, c, "tag", errors)
        source = _key_value(cell(4, c), 4, c, "source", errors)
        selector = pointer = None
        if name is not None and not name:
            errors.append(f"R2C{c}: the column name is empty")
        if tag is not None:
            selector = TAG_WORDS.get(tag)
            if selector is None:
                near = [w for w in TAG_WORDS if w.lower() == tag.lower()]
                hint = f" (tag words are case-sensitive: use '{near[0]}')" if near else ""
                errors.append(f"R3C{c}: unknown tag selector {tag!r}{hint}")
            elif selector not in LOGTABLE_SELECTORS:
                errors.append(f"R3C{c}: a log table cannot hold a '{tag}' column")
                selector = None
        if source is not None:
            try:
                pointer = parse_cell(source)
            except ValueError:
                errors.append(f"R4C{c}: invalid source {source!r}; expected '<n>' (a PID) "
                              "or '<n> VPID'")
        if is_event:
            if selector is not None and selector != EVENT_SELECTOR:
                errors.append(f"R3C{c}: the column marked 'event' must have 'tag: dateTime'")
            if pointer is not None and pointer.kind != "PID":
                errors.append(f"R4C{c}: the column marked 'event' must have a PID source, "
                              "not a VPID")
        if name and selector and pointer:
            cols.append(ColumnSpec(name, pointer, selector, event=is_event))

    marked = [c for c in range(1, width + 1)
              if isinstance(cell(1, c), str) and cell(1, c).strip() == "event"]
    if not marked:
        errors.append("no column is marked 'event' in row 1; the server rejects this as "
                      "'Missing event source column'")
    elif len(marked) > 1:
        errors.append(f"row 1 marks {len(marked)} columns 'event' "
                      f"({', '.join(f'C{c}' for c in marked)}); exactly one is allowed")
    if errors:
        raise ValueError("the file's header is invalid:\n  - " + "\n  - ".join(errors))
    notes = []
    if data_rows:
        notes.append(f"{data_rows} data row(s) from row 5 on are ignored; a new table starts empty")
    return Header(cols, data_rows=data_rows, sheet=sheet, notes=notes)


def read_header(path: str | Path) -> Header:
    """Read an `onping-logtable-export` XLSX (or a hand-made one in the same shape)."""
    try:
        from openpyxl import load_workbook
    except ImportError:  # pragma: no cover - uv installs it from the script header
        raise ValueError("--from needs openpyxl (declared in the script's uv header)") from None
    p = Path(path)
    try:
        data = p.read_bytes()
    except OSError as e:
        raise ValueError(f"cannot read {p}: {e.strerror or e}") from None
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:  # openpyxl raises several unrelated types for a bad file
        raise ValueError(f"{p} is not a readable XLSX workbook: {e}") from None
    try:
        sheet = "Sheet1" if "Sheet1" in wb.sheetnames else wb.sheetnames[0]
        rows: list[list[Any]] = []
        data_rows = 0
        for i, row in enumerate(wb[sheet].iter_rows(values_only=True)):
            if i < 4:
                rows.append(list(row))
            elif any(not _blank(v) for v in row):
                data_rows += 1
    finally:
        wb.close()
    header = parse_header_rows(rows, data_rows=data_rows, sheet=sheet)
    if sheet != "Sheet1":
        header.notes.append(f"read sheet {sheet!r} (the file has no 'Sheet1')")
    return header


# ------------------------------------------------------------------ schema

def build_schema(columns: list[ColumnSpec]) -> tuple[dict, list[str]]:
    """The version-1 schema holding `columns` in order. Returns (schema, column UUIDs in order).

    Exactly one column must be the event column (a PID with LastUpdateDateTime). The
    others are added with `logtable.add_column`, which refuses empty or repeated names,
    a repeated key and selector, and selectors a log table cannot hold.
    """
    events = [c for c in columns if c.event]
    if len(events) != 1:
        raise ValueError(f"expected exactly one event column, found {len(events)}")
    ev = events[0]
    if ev.pointer.kind != "PID" or ev.selector != EVENT_SELECTOR:
        raise ValueError("the event column must be LastUpdateDateTime of a PID")
    sch, event_cid = lt.empty_schema(ev.pointer.value, ev.name)
    cids = []
    for i, col in enumerate(columns):
        if col.event:
            cids.append(event_cid)
            continue
        try:
            cids.append(lt.add_column(sch, col.pointer, col.name, col.selector, at=i))
        except ValueError as e:
            raise ValueError(f"column {i + 1} ({col.name!r}): {e}") from None
    errors = lt.validate_schema(sch)
    if errors:
        raise ValueError("the version-1 schema is invalid: " + "; ".join(errors))
    return sch, cids
