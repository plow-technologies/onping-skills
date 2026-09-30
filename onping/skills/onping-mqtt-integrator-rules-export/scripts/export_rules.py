# /// script
# requires-python = ">=3.10"
# dependencies = ["requests", "openpyxl"]
# ///
"""Download one Lumberjack's mqtt-json-integrator generation rules.

Routes: GET /mqtt/json/integrator/{serial}/rules/export/{filename}   -> XLSX
        GET /mqtt/json/integrator/{serial}/generation/rules          -> JSON

Read-only on the server: both are GETs and neither mutates state.

TWO READ PATHS, DIFFERENT FIDELITY.
  default   the XLSX export — the round-trip format consumed by
            onping-mqtt-integrator-rules-import. 12 columns, one row per
            (location rule, PID rule) pair, plus location-only rows.
  --json    the JSON rule list — carries the LocationRuleIdentifier and
            PidRuleIdentifier values that the spreadsheet OMITS. Use this when
            you need to target a specific rule for an incremental update, since
            the sheet cannot express rule ids at all.

THE {filename} PATH SEGMENT IS IGNORED by the handler (it binds it to `_`). It
exists only so a browser names the download. Passing anything non-empty works;
this script defaults to "rules.xlsx" and uses --out for the real local path.

WHY YOU ALMOST ALWAYS WANT --summary FIRST. Rules import REPLACES THE ENTIRE
rule set and renumbers every id from row order, so the safe workflow is
export -> edit -> re-import with every rule you intend to keep still present.
Reading the sheet before editing it is the only way to know what "everything"
currently is.

Source of truth (re-verify if these drift):
  - Handlers: onping/Handler/MqttJsonIntegrator/Service.hs (XLSX export),
              :132 (JSON list)
  - Sheet:    onping/Handler/MqttJsonIntegrator/Rules/ImportExport.hs
  - Schema:   _mqtt_integrator_routes/rules_sheet.py (RULES_HEADERS)
"""

from __future__ import annotations

import argparse
import io
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))  # the skills root
from _mqtt_integrator_routes import rules_sheet
from _mqtt_integrator_routes.integrator_http import _fail, get_bytes, get_json
from _mqtt_integrator_routes.routes import endpoint

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - uv installs the dep
    print("openpyxl is required (declared in the uv script header)", file=sys.stderr)
    sys.exit(1)


def read_sheet(xlsx: bytes) -> tuple[list, list[list]]:
    """Parse the exported workbook into (header_row, data_rows)."""
    try:
        wb = load_workbook(io.BytesIO(xlsx), data_only=True)
    except Exception as e:  # openpyxl raises a zoo of exception types
        _fail(f"Could not read the returned workbook: {e}")
    if rules_sheet.SHEET_NAME not in wb.sheetnames:
        _fail(
            f"Expected a {rules_sheet.SHEET_NAME!r} worksheet, found "
            f"{wb.sheetnames}. The export format may have changed."
        )
    ws = wb[rules_sheet.SHEET_NAME]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    if not rows:
        return [], []
    header = [c for c in rows[rules_sheet.HEADER_ROW - 1]]
    data = rows[rules_sheet.FIRST_DATA_ROW - 1 :]

    # The server filters blank cells on read; mirror that so trailing styled
    # rows from the xlsx template do not read as empty rules.
    def is_blank(row: list) -> bool:
        return all(
            v is None or (isinstance(v, str) and not v.strip()) for v in row
        )

    while data and is_blank(data[-1]):
        data.pop()
    return header, data


def main() -> None:
    p = argparse.ArgumentParser(
        description=(
            "Download one Lumberjack's mqtt-json-integrator generation rules as "
            "XLSX (default) or JSON. Read-only."
        ),
    )
    p.add_argument("access_token", help="OnPing OAuth bearer token")
    p.add_argument("serial", help="LJSerial of the Lumberjack running the integrator")
    p.add_argument(
        "--out",
        help="local path to write (default: integrator-rules-<serial>.xlsx, or "
        "-<serial>.json with --json)",
    )
    p.add_argument(
        "--json",
        action="store_true",
        help="fetch the JSON rule list instead of the XLSX — includes the rule "
        "IDENTIFIERS that the spreadsheet omits",
    )
    p.add_argument(
        "--summary",
        action="store_true",
        help="print a per-rule summary of what was downloaded (implies reading "
        "the sheet locally)",
    )
    p.add_argument(
        "--validate",
        action="store_true",
        help="check the downloaded sheet against the expected schema and report "
        "any row the import would reject",
    )
    p.add_argument(
        "--filename",
        default="rules.xlsx",
        help="the {filename} path segment; IGNORED by the server "
        "(default: rules.xlsx)",
    )
    p.add_argument(
        "--stdout",
        action="store_true",
        help="with --json, print to stdout instead of writing a file",
    )
    args = p.parse_args()

    # ── JSON path ───────────────────────────────────────────────────────────
    if args.json:
        rules = get_json(
            args.access_token, endpoint("get_rules"), serial=args.serial
        )
        if not isinstance(rules, list):
            _fail(f"Expected a rule array, got: {json.dumps(rules)[:400]}")
        text = json.dumps(rules, indent=2, sort_keys=True)
        if args.stdout:
            print(text)
        else:
            out = Path(args.out or f"integrator-rules-{args.serial}.json")
            out.write_text(text + "\n", encoding="utf-8")
            print(f"Wrote {len(rules)} rule(s) to {out}")
        if args.summary:
            print(f"\n{len(rules)} location rule(s):", file=sys.stderr)
            for r in rules:
                loc = r.get("generationRuleLocation", {})
                pids = r.get("generationRulePids", []) or []
                rid = (loc.get("locationRuleLocationRuleId") or {}).get(
                    "unLocationRuleIdentifier"
                )
                match = (loc.get("locationRuleLocationId") or {}).get("unMqttRuleText")
                print(
                    f"  id={rid}  selector={loc.get('locationRuleSelectorName')!r}  "
                    f"match={match!r}  pids={len(pids)}",
                    file=sys.stderr,
                )
        return

    # ── XLSX path ───────────────────────────────────────────────────────────
    xlsx = get_bytes(
        args.access_token,
        endpoint("export_rules"),
        serial=args.serial,
        filename=args.filename,
    )
    out = Path(args.out or f"integrator-rules-{args.serial}.xlsx")
    out.write_bytes(xlsx)
    print(f"Wrote {len(xlsx)} bytes to {out}")

    if not (args.summary or args.validate):
        print(
            "\nThis sheet is the round-trip format for "
            "onping-mqtt-integrator-rules-import, which REPLACES ALL RULES. Keep "
            "every rule you intend to retain."
        )
        return

    header, rows = read_sheet(xlsx)
    expected = rules_sheet.RULES_HEADERS
    got = [str(h).strip() if h is not None else "" for h in header][: len(expected)]
    if got != expected:
        print(
            f"\nWARNING: header row does not match the expected schema.\n"
            f"  expected: {expected}\n  got:      {got}\n"
            f"The export format may have changed; re-verify "
            f"Rules/ImportExport.hs against rules_sheet.RULES_HEADERS.",
            file=sys.stderr,
        )

    location_only = 0
    pid_rows = 0
    errors: list[str] = []
    blocks: list[tuple[tuple, int]] = []  # (location key, pid count) in row order

    for offset, row in enumerate(rows):
        row_number = rules_sheet.FIRST_DATA_ROW + offset
        row_errors, is_loc_only = rules_sheet.validate_row(
            row, flavor="rules", row_number=row_number
        )
        errors.extend(row_errors)
        if is_loc_only:
            location_only += 1
        else:
            pid_rows += 1
        key = tuple(
            (str(v).strip() if v is not None else "") for v in list(row)[:3]
        )
        if blocks and blocks[-1][0] == key:
            blocks[-1] = (key, blocks[-1][1] + (0 if is_loc_only else 1))
        else:
            blocks.append((key, 0 if is_loc_only else 1))

    if args.summary:
        print(
            f"\n{len(rows)} data row(s): {pid_rows} with a PID rule, "
            f"{location_only} location-only."
        )
        print(f"{len(blocks)} contiguous location block(s) — this is how the "
              f"import will group them:\n")
        for (selector, name, match), count in blocks:
            print(f"  selector={selector!r}")
            print(f"      name={name!r}  match={match!r}  pid rules={count}")

        seen: dict[tuple, int] = {}
        for key, _ in blocks:
            seen[key] = seen.get(key, 0) + 1
        repeated = {k: v for k, v in seen.items() if v > 1}
        if repeated:
            print(
                f"\nWARNING: {len(repeated)} location appears in more than one "
                f"NON-ADJACENT block. The import groups only CONSECUTIVE rows "
                f"with identical columns 1-3, so each block becomes a SEPARATE "
                f"location rule. Sort the sheet so every location's rows are "
                f"contiguous.",
                file=sys.stderr,
            )
            for key, count in repeated.items():
                print(f"  {count}x  selector={key[0]!r} match={key[2]!r}", file=sys.stderr)

    if args.validate:
        if errors:
            print(f"\n{len(errors)} schema problem(s) in the DOWNLOADED sheet:",
                  file=sys.stderr)
            for e in errors:
                print(f"  - {e}", file=sys.stderr)
            print(
                "\nThese came straight from the server, so they indicate either a "
                "format drift or rules the server itself would no longer accept.",
                file=sys.stderr,
            )
            sys.exit(1)
        print("\nValidation: every row matches the expected rules-sheet schema.")


if __name__ == "__main__":
    main()
